import os
import re
import sqlite3
import secrets
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from functools import wraps
from pathlib import Path
from zoneinfo import ZoneInfo
from urllib.parse import urlencode
import requests
from flask import Flask, Response, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook

APP_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("DATABASE_PATH", "/data/ledger.sqlite3"))
DB_PATH.parent.mkdir(parents=True, exist_ok=True)
app = Flask(__name__)
app.secret_key = os.environ.get("SESSION_SECRET", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024

def db_connect():
    db = sqlite3.connect(DB_PATH, timeout=30)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.execute("PRAGMA journal_mode=WAL")
    db.executescript("""
    CREATE TABLE IF NOT EXISTS periods(id INTEGER PRIMARY KEY, started TEXT NOT NULL, payday TEXT, closed INTEGER NOT NULL DEFAULT 0);
    CREATE TABLE IF NOT EXISTS items(
      order_id TEXT NOT NULL, line_no INTEGER NOT NULL, period_id INTEGER NOT NULL,
      ordered_at TEXT, shop_id TEXT, order_label TEXT, product_id TEXT, variant_id TEXT,
      title TEXT, variant TEXT, sku TEXT, quantity INTEGER, status TEXT,
      retail INTEGER, production INTEGER, shipping INTEGER,
      owner TEXT NOT NULL DEFAULT 'Unassigned', owner_share INTEGER, rykeen_gross INTEGER,
      rykeen_tax INTEGER, rykeen_net INTEGER, updated_at TEXT,
      PRIMARY KEY(order_id,line_no), FOREIGN KEY(period_id) REFERENCES periods(id));
    CREATE TABLE IF NOT EXISTS products(product_id TEXT PRIMARY KEY, title TEXT NOT NULL, owner TEXT NOT NULL DEFAULT 'Unassigned');
    CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
    """)
    if not db.execute("SELECT 1 FROM periods LIMIT 1").fetchone():
        db.execute("INSERT INTO periods(started) VALUES (?)", (datetime.now(timezone.utc).isoformat(),))
    db.execute("INSERT OR IGNORE INTO settings(key,value) SELECT 'next_session_start_date',value FROM settings WHERE key='next_cutoff_date'")
    db.execute("INSERT OR IGNORE INTO settings(key,value) VALUES('next_session_start_date','2026-09-20')")
    db.commit()
    return db

def cents(value):
    if value is None:
        return None
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))

def share(amount, pct):
    return int((Decimal(amount) * Decimal(pct) / 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))

def payout(retail, production):
    profit = max(0, retail - production)
    rykeen = share(profit, 25)
    tax = share(rykeen, 25)
    return profit, profit - rykeen, rykeen, tax, rykeen - tax

def order_local_date(value):
    if not value:
        return None
    try:
        if len(value) == 10:
            return date.fromisoformat(value)
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(ZoneInfo("America/New_York")).date()
    except (TypeError, ValueError):
        return None

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login", next=request.path))
        return fn(*args, **kwargs)
    return wrapped

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        expected = os.environ.get("APP_PASSWORD", "")
        if not expected:
            return "APP_PASSWORD is not configured. Set it in Railway Variables.", 503
        if secrets.compare_digest(request.form.get("password", ""), expected):
            session["logged_in"] = True
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("That password did not match.", "error")
    return render_template("login.html")

@app.post("/logout")
@login_required
def logout():
    session.clear()
    return redirect(url_for("login"))

def current_period(db):
    row = db.execute("SELECT id FROM periods WHERE closed=0 ORDER BY id DESC LIMIT 1").fetchone()
    return row["id"]

def printify_get(path):
    token = os.environ.get("PRINTIFY_API_TOKEN", "").strip()
    if not token:
        raise RuntimeError("Set PRINTIFY_API_TOKEN in Railway Variables before syncing.")
    response = requests.get("https://api.printify.com/v1/" + path, headers={
        "Authorization": "Bearer " + token, "User-Agent": "PrintifyAudit/1.0",
        "Accept": "application/json"}, timeout=45)
    response.raise_for_status()
    return response.json()

def paged(path):
    page = 1
    while True:
        payload = printify_get(f"{path}?limit=10&page={page}")
        rows = payload.get("data") or []
        yield from rows
        last = int(payload.get("last_page") or (page if len(rows) < 10 else page + 1))
        if not rows or page >= last:
            break
        page += 1

def title_owner(title):
    parts = re.split(r"\s*[-–—]\s*", title, maxsplit=1)
    prefix = parts[0].strip()
    # Only treat a single handle-like title prefix as an owner. Long product
    # descriptions before a dash are not owner names.
    return prefix.title() if len(parts) > 1 and re.fullmatch(r"[A-Za-z0-9_]+", prefix) else "Unassigned"

def sync_shop_data(db):
    shops = printify_get("shops.json")
    product_count = order_count = 0
    for shop in shops:
        shop_id = str(shop["id"])
        for product in paged(f"shops/{shop_id}/products.json"):
            product_id = str(product.get("id") or "")
            title = str(product.get("title") or "Untitled product")
            if not product_id:
                continue
            prior = db.execute("SELECT owner FROM products WHERE product_id=?", (product_id,)).fetchone()
            owner = prior["owner"] if prior else title_owner(title)
            db.execute("""INSERT INTO products(product_id,title,owner) VALUES(?,?,?)
                ON CONFLICT(product_id) DO UPDATE SET title=excluded.title""",
                (product_id, title, owner))
            product_count += 1
        db.commit()
        for order in paged(f"shops/{shop_id}/orders.json"):
            order_id = str(order.get("id") or "")
            if not order_id:
                continue
            current = current_period(db)
            for line_no, line in enumerate(order.get("line_items") or []):
                old = db.execute("SELECT period_id,owner FROM items WHERE order_id=? AND line_no=?", (order_id,line_no)).fetchone()
                if old and db.execute("SELECT closed FROM periods WHERE id=?", (old["period_id"],)).fetchone()["closed"]:
                    continue
                meta = line.get("metadata") or {}
                product_id = str(line.get("product_id") or "")
                title = str(meta.get("title") or "Unknown item")
                product = db.execute("SELECT owner FROM products WHERE product_id=?", (product_id,)).fetchone()
                owner = old["owner"] if old else (product["owner"] if product else title_owner(title))
                if product_id:
                    db.execute("""INSERT INTO products(product_id,title,owner) VALUES(?,?,?)
                        ON CONFLICT(product_id) DO UPDATE SET title=excluded.title""",(product_id,title,owner))
                qty = max(1, int(line.get("quantity") or 1))
                retail = cents(meta.get("price"))
                retail = retail * qty if retail is not None else None
                production = int(line.get("cost") or 0)
                shipping = int(line.get("shipping_cost") or 0)
                status = str(line.get("status") or order.get("status") or "unknown")
                if retail is not None:
                    profit, owner_cut, rykeen, tax, net = payout(retail, production)
                else:
                    owner_cut = rykeen = tax = net = None
                db.execute("""INSERT INTO items VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(order_id,line_no) DO UPDATE SET
                    ordered_at=excluded.ordered_at,status=excluded.status,
                    retail=COALESCE(excluded.retail,items.retail),production=excluded.production,
                    shipping=excluded.shipping,owner_share=excluded.owner_share,
                    rykeen_gross=excluded.rykeen_gross,rykeen_tax=excluded.rykeen_tax,
                    rykeen_net=excluded.rykeen_net,updated_at=excluded.updated_at""",
                    (order_id,line_no,old["period_id"] if old else current,order.get("created_at"),
                     shop_id,str((order.get("metadata") or {}).get("shop_order_label") or order.get("app_order_id") or order_id),
                     product_id,str(line.get("variant_id") or ""),title,str(meta.get("variant_label") or ""),
                     str(meta.get("sku") or ""),qty,status,retail,production,shipping,owner,owner_cut,
                     rykeen,tax,net,datetime.now(timezone.utc).isoformat()))
                order_count += 1
            db.commit()
    return product_count, order_count

@app.route("/")
@login_required
def dashboard():
    db = db_connect()
    periods = db.execute("SELECT * FROM periods ORDER BY id DESC").fetchall()
    active = current_period(db)
    chosen = request.args.get("period", type=int) or active
    if not any(p["id"] == chosen for p in periods):
        chosen = active
    default_start = db.execute("SELECT value FROM settings WHERE key='next_session_start_date'").fetchone()["value"]
    session_start = request.args.get("session_start", default_start)
    try:
        start_day = date.fromisoformat(session_start)
        if start_day.isoformat() != session_start:
            raise ValueError()
    except ValueError:
        session_start = default_start
        start_day = date.fromisoformat(session_start)
    rows = db.execute("SELECT * FROM items WHERE period_id=?", (chosen,)).fetchall()
    if chosen == active:
        rows = [r for r in rows if order_local_date(r["ordered_at"]) is not None and order_local_date(r["ordered_at"]) >= start_day]
    owner_order = lambda r: (r["owner"].lower() == "unassigned", r["owner"].casefold(), r["title"].casefold(), r["ordered_at"] or "")
    rows.sort(key=owner_order)
    sums = {"retail":0,"production":0,"owner_share":0,"rykeen_tax":0,"rykeen_net":0}
    for r in rows:
        if r["retail"] is None or r["status"].lower() in ("canceled","cancelled","refunded"):
            continue
        for key in sums:
            sums[key] += r[key] or 0
    products = db.execute("SELECT count(*) n FROM products").fetchone()["n"]
    db.close()
    return render_template("dashboard.html", rows=rows, periods=periods, chosen=chosen, active=active,
                           sums=sums, products=products, session_start=session_start)

@app.post("/sync")
@login_required
def sync():
    db = db_connect()
    try:
        products, orders = sync_shop_data(db)
        flash(f"Sync finished: {products} products scanned and {orders} order lines checked.", "success")
    except Exception as exc:
        flash(f"Sync failed: {exc}", "error")
    finally:
        db.close()
    return redirect(url_for("dashboard"))

@app.route("/products", methods=["GET", "POST"])
@login_required
def products():
    db = db_connect()
    if request.method == "POST":
        owner = request.form.get("owner","").strip()
        selected = request.form.getlist("product_id")
        if owner and selected:
            db.executemany("UPDATE products SET owner=? WHERE product_id=?", [(owner,pid) for pid in selected])
            active = current_period(db)
            for pid in selected:
                for row in db.execute("SELECT order_id,line_no,retail,production FROM items WHERE product_id=? AND period_id=?",(pid,active)).fetchall():
                    if row["retail"] is not None:
                        _, owner_cut, rykeen, tax, net = payout(row["retail"],row["production"])
                        db.execute("UPDATE items SET owner=?,owner_share=?,rykeen_gross=?,rykeen_tax=?,rykeen_net=? WHERE order_id=? AND line_no=?",(owner,owner_cut,rykeen,tax,net,row["order_id"],row["line_no"]))
                    else:
                        db.execute("UPDATE items SET owner=? WHERE order_id=? AND line_no=?",(owner,row["order_id"],row["line_no"]))
            db.commit()
            flash(f"Assigned {len(selected)} product(s) to {owner}.", "success")
        db.close()
        return redirect(url_for("products"))
    query = request.args.get("q","").strip()
    if query:
        result = db.execute("SELECT * FROM products WHERE title LIKE ? OR owner LIKE ? ORDER BY title COLLATE NOCASE", (f"%{query}%",f"%{query}%")).fetchall()
    else:
        result = db.execute("SELECT * FROM products ORDER BY title COLLATE NOCASE").fetchall()
    db.close()
    return render_template("products.html", products=result, query=query)

def create_workbook(db, period_id, path, session_start_date=None):
    rows = db.execute("SELECT * FROM items WHERE period_id=? ORDER BY title,variant,ordered_at",(period_id,)).fetchall()
    if session_start_date:
        start_day = date.fromisoformat(session_start_date)
        rows = [r for r in rows if order_local_date(r["ordered_at"]) is not None and order_local_date(r["ordered_at"]) >= start_day]
    wb = Workbook()
    summary = wb.active
    summary.title = "Payout Summary"
    summary.append(["Product","Owner","Qty sold","Retail","Production cost","Profit","Owner payout (75%)","Rykeen gross (25%)","Rykeen tax reserve","Rykeen net"])
    totals = {}
    for r in rows:
        if r["retail"] is None or r["status"].lower() in ("canceled","cancelled","refunded"):
            continue
        key=(r["title"],r["owner"])
        t=totals.setdefault(key,[0,0,0,0,0,0,0])
        profit=max(0,r["retail"]-r["production"])
        for i,v in enumerate((r["quantity"],r["retail"],r["production"],profit,r["owner_share"],r["rykeen_gross"],r["rykeen_tax"])):
            t[i]+=v or 0
    for (title,owner),v in sorted(totals.items()):
        summary.append([title,owner,v[0],v[1]/100,v[2]/100,v[3]/100,v[4]/100,v[5]/100,v[6]/100,(v[5]-v[6])/100])
    detail=wb.create_sheet("Sale Details")
    detail.append(["Sale date","Order","Product","Variant","Qty","Status","Owner","Retail","Production cost","Shipping cost","Profit","Owner payout","Rykeen gross","Rykeen reserve","Rykeen net"])
    for r in rows:
        valid=r["retail"] is not None and r["status"].lower() not in ("canceled","cancelled","refunded")
        profit=max(0,r["retail"]-r["production"]) if valid else 0
        detail.append([r["ordered_at"],r["order_label"],r["title"],r["variant"],r["quantity"],r["status"],r["owner"],
            (r["retail"] or 0)/100 if valid else 0,r["production"]/100,r["shipping"]/100,profit/100,
            (r["owner_share"] or 0)/100 if valid else 0,(r["rykeen_gross"] or 0)/100 if valid else 0,
            (r["rykeen_tax"] or 0)/100 if valid else 0,(r["rykeen_net"] or 0)/100 if valid else 0])
    for sheet in (summary,detail):
        sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:
            letter=col[0].column_letter
            sheet.column_dimensions[letter].width=min(42,max(12,max(len(str(c.value or "")) for c in col)+2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.column >= (4 if sheet is summary else 8) and isinstance(cell.value,(int,float)):
                    cell.number_format='"$"#,##0.00'
    for cell in summary["C"][1:]: cell.number_format="0"
    for cell in detail["E"][1:]: cell.number_format="0"
    wb.save(path)

@app.post("/close-period")
@login_required
def close_period():
    db = db_connect()
    period_id = current_period(db)
    session_start_date = request.form.get("session_start_date","").strip()
    try:
        start_day = date.fromisoformat(session_start_date)
        if start_day.isoformat() != session_start_date:
            raise ValueError()
    except ValueError:
        db.close()
        flash("Choose a valid session start date.", "error")
        return redirect(url_for("dashboard"))
    open_rows = db.execute("SELECT * FROM items WHERE period_id=?",(period_id,)).fetchall()
    included = [r for r in open_rows if order_local_date(r["ordered_at"]) is not None and order_local_date(r["ordered_at"]) >= start_day]
    missing = sum(1 for r in included if r["owner"] == "Unassigned" or r["retail"] is None)
    if missing:
        db.close()
        flash(f"Assign owners and enter actual retail for {missing} line(s) before closing.", "error")
        return redirect(url_for("dashboard"))
    from tempfile import NamedTemporaryFile
    temp=NamedTemporaryFile(prefix=f"Printify-Payout-{period_id}-",suffix=".xlsx",delete=False)
    temp.close()
    try:
        create_workbook(db,period_id,temp.name,session_start_date)
        now=datetime.now(timezone.utc).isoformat()
        db.execute("UPDATE periods SET closed=1,payday=? WHERE id=?",(now,period_id))
        db.execute("INSERT INTO periods(started) VALUES (?)",(now,))
        today=datetime.now(ZoneInfo("America/New_York")).date().isoformat()
        db.execute("UPDATE settings SET value=? WHERE key='next_session_start_date'",(today,))
        db.commit()
    except Exception:
        db.close()
        os.unlink(temp.name)
        raise
    db.close()
    response=send_file(temp.name,as_attachment=True,download_name=f"Printify-Payout-Period-{period_id}.xlsx",
                       mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.call_on_close(lambda: os.path.exists(temp.name) and os.unlink(temp.name))
    return response

@app.get("/health")
def health():
    return {"status":"ok"}

@app.errorhandler(413)
def too_large(_):
    return "Request too large.",413

if __name__ == "__main__":
    app.run(host="0.0.0.0",port=int(os.environ.get("PORT","8080")))
